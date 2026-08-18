from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter



def generate_outstanding_report_excel(
    data,
    filename
):

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Outstanding Fees"



    # ==========================
    # TITLE
    # ==========================

    sheet["A1"] = (
        "CAMPUSPAY ENTERPRISE"
    )

    sheet["A2"] = (
        "Outstanding Fees Report"
    )


    sheet["A1"].font = Font(
        bold=True,
        size=16
    )


    sheet["A2"].font = Font(
        bold=True,
        size=12
    )


    # ==========================
    # SUMMARY
    # ==========================


    sheet["A4"] = "Total Expected"

    sheet["B4"] = (
        data["total_expected"]
    )


    sheet["A5"] = "Total Paid"

    sheet["B5"] = (
        data["total_paid"]
    )


    sheet["A6"] = "Total Outstanding"

    sheet["B6"] = (
        data["total_outstanding"]
    )



    # ==========================
    # TABLE HEADER
    # ==========================


    start_row = 8


    headers = [

        "Student ID",

        "Full Name",

        "Department",

        "Level",

        "Expected Fee",

        "Amount Paid",

        "Outstanding Balance",

        "Status",

    ]


    for col, header in enumerate(
        headers,
        start=1
    ):

        cell = sheet.cell(
            row=start_row,
            column=col
        )

        cell.value = header


        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )


        cell.fill = PatternFill(
            "solid",
            fgColor="1F4E78"
        )


        cell.alignment = Alignment(
            horizontal="center"
        )



    # ==========================
    # DATA ROWS
    # ==========================


    row = start_row + 1


    for item in data["report"]:


        values = [

            item["member"].student_id,

            item["member"].full_name,

            str(
                item["member"].department
            ),

            item["member"].level,

            item["expected"],

            item["paid"],

            item["balance"],

            item["status"],

        ]


        for col, value in enumerate(
            values,
            start=1
        ):

            sheet.cell(
                row=row,
                column=col
            ).value = value


        row += 1



    # ==========================
    # CURRENCY FORMAT
    # ==========================


    for r in range(
        start_row + 1,
        row
    ):

        for c in [
            5,
            6,
            7
        ]:

            sheet.cell(
                r,
                c
            ).number_format = (
                '"GH₵" #,##0.00'
            )



    # ==========================
    # COLUMN WIDTH
    # ==========================


    widths = {

        "A":15,
        "B":25,
        "C":25,
        "D":12,
        "E":18,
        "F":18,
        "G":22,
        "H":15,

    }


    for column, width in widths.items():

        sheet.column_dimensions[
            column
        ].width = width



    # Freeze table header

    sheet.freeze_panes = "A9"
    
    sheet.auto_filter.ref = (
    f"A8:H{row-1}"
    )



    workbook.save(
        filename
    )